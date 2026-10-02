import os
import re
import shutil

from openpyxl import load_workbook

from services.app_paths import AppPaths
from services.safe_storage import SafeStorage


class VendorMasterService:

    BASE_DIR = str(AppPaths.install_dir())

    DATA_DIR = str(AppPaths.data_dir())

    VENDOR_FILE = os.path.join(
        DATA_DIR,
        "vendor_master.json"
    )

    CATEGORY_FILE = os.path.join(
        DATA_DIR,
        "vendor_categories.json"
    )

    CATEGORY_IMAGE_DIR = os.path.join(
        DATA_DIR,
        "category_images"
    )

    CATEGORY_IMAGE_FILE = os.path.join(
        DATA_DIR,
        "category_images.json"
    )

    DEFAULT_CATEGORIES = [
        "滅菌包",
        "不織布",
        "百片紗布",
        "棒裝紗布",
        "棉棒",
        "尿袋抽痰包",
        "女帽隔離衣",
        "口罩",
    ]

    # =====================================================
    # Folder
    # =====================================================

    @staticmethod
    def ensure_data_folder():

        os.makedirs(
            VendorMasterService.DATA_DIR,
            exist_ok=True
        )

        os.makedirs(
            VendorMasterService.CATEGORY_IMAGE_DIR,
            exist_ok=True
        )

    # =====================================================
    # Vendor master
    # =====================================================

    @staticmethod
    def load_vendors():

        VendorMasterService.ensure_data_folder()

        data = SafeStorage.safe_read_json(
            VendorMasterService.VENDOR_FILE,
            default=[]
        )

        if isinstance(data, list):
            return data

        return []

    @staticmethod
    def save_vendors(vendors):
        from services import security
        security.require_write_access("儲存廠商主檔")

        VendorMasterService.ensure_data_folder()

        SafeStorage.atomic_write_json(
            VendorMasterService.VENDOR_FILE,
            vendors,
            indent=4
        )

    # =====================================================
    # Vendor edit
    # =====================================================

    @staticmethod
    def update_vendor(
        old_name,
        new_name,
        phone,
        fax,
        address
    ):

        vendors = (
            VendorMasterService.load_vendors()
        )

        for vendor in vendors:

            if str(
                vendor.get(
                    "name",
                    ""
                )
            ).strip() == str(
                old_name
            ).strip():

                vendor["name"] = (
                    str(new_name).strip()
                )

                vendor["phone"] = (
                    str(phone).strip()
                )

                vendor["fax"] = (
                    str(fax).strip()
                )

                vendor["address"] = (
                    str(address).strip()
                )

                break

        VendorMasterService.save_vendors(
            vendors
        )

        return vendors

    # =====================================================
    # Categories
    # =====================================================

    @staticmethod
    def load_categories():

        VendorMasterService.ensure_data_folder()

        categories = list(
            VendorMasterService.DEFAULT_CATEGORIES
        )

        if os.path.exists(
            VendorMasterService.CATEGORY_FILE
        ):

            saved = SafeStorage.safe_read_json(
                VendorMasterService.CATEGORY_FILE,
                default=[]
            )

            if isinstance(saved, list):

                for category in saved:

                    category = str(
                        category
                    ).strip()

                    if (
                        category
                        and category not in categories
                    ):

                        categories.append(
                            category
                        )

        return categories

    @staticmethod
    def save_categories(categories):
        from services import security
        security.require_write_access("儲存廠商分類")

        VendorMasterService.ensure_data_folder()

        result = []

        for category in categories:

            category = str(
                category
            ).strip()

            if (
                category
                and category not in result
            ):

                result.append(
                    category
                )

        SafeStorage.atomic_write_json(
            VendorMasterService.CATEGORY_FILE,
            result,
            indent=4
        )

    @staticmethod
    def add_category(category):

        category = str(
            category
        ).strip()

        if not category:
            return False

        categories = (
            VendorMasterService.load_categories()
        )

        if category in categories:
            return False

        categories.append(
            category
        )

        VendorMasterService.save_categories(
            categories
        )

        return True

    # =====================================================
    # Category image
    # =====================================================

    @staticmethod
    def load_category_images():

        VendorMasterService.ensure_data_folder()

        if not os.path.exists(
            VendorMasterService.CATEGORY_IMAGE_FILE
        ):

            return {}

        data = SafeStorage.safe_read_json(
            VendorMasterService.CATEGORY_IMAGE_FILE,
            default={}
        )

        if isinstance(data, dict):
            return data

        return {}

    @staticmethod
    def save_category_images(data):
        from services import security
        security.require_write_access("儲存廠商分類圖片")

        VendorMasterService.ensure_data_folder()

        SafeStorage.atomic_write_json(
            VendorMasterService.CATEGORY_IMAGE_FILE,
            data,
            indent=4
        )

    @staticmethod
    def set_category_image(
        category,
        source_file
    ):

        VendorMasterService.ensure_data_folder()

        category = str(
            category
        ).strip()

        if not category:
            return None

        extension = os.path.splitext(
            source_file
        )[1].lower()

        if extension not in [
            ".png",
            ".jpg",
            ".jpeg",
            ".webp",
        ]:

            return None

        safe_name = re.sub(
            r'[\\/:*?"<>|]',
            "_",
            category
        )

        destination = os.path.join(
            VendorMasterService.CATEGORY_IMAGE_DIR,
            safe_name + extension
        )

        shutil.copy2(
            source_file,
            destination
        )

        image_map = (
            VendorMasterService
            .load_category_images()
        )

        image_map[
            category
        ] = destination

        VendorMasterService.save_category_images(
            image_map
        )

        return destination

    # =====================================================
    # Text helpers
    # =====================================================

    @staticmethod
    def clean_text(value):

        if value is None:
            return ""

        return str(value).strip()

    # =====================================================
    # 重點：
    # 同一格可能是：
    #
    # 賣方：ABC公司 電話：123 傳真：456 地址：xxxx
    #
    # 必須只抓標籤到下一個標籤之間
    # =====================================================

    @staticmethod
    def extract_field(
        text,
        labels
    ):

        text = (
            VendorMasterService
            .clean_text(
                text
            )
        )

        if not text:
            return ""

        all_labels = [
            "賣方",
            "卖方",
            "電話",
            "电话",
            "TEL",
            "Tel",
            "傳真",
            "传真",
            "FAX",
            "Fax",
            "地址",
            "公司地址",
            "工廠地址",
            "厂址",
            "買方",
            "买方",
        ]

        next_label_pattern = "|".join(
            re.escape(label)
            for label in all_labels
        )

        for label in labels:

            pattern = (
                rf"{re.escape(label)}"
                rf"\s*[:：]\s*"
                rf"(.*?)"
                rf"(?=\s*(?:{next_label_pattern})\s*[:：]|$)"
            )

            match = re.search(
                pattern,
                text,
                flags=re.IGNORECASE
            )

            if match:

                return (
                    match.group(1)
                    .strip()
                )

        return ""

    @staticmethod
    def get_next_value(
        values,
        index
    ):

        for next_index in range(
            index + 1,
            len(values)
        ):

            value = (
                VendorMasterService
                .clean_text(
                    values[
                        next_index
                    ]
                )
            )

            if value:
                return value

        return ""

    # =====================================================
    # Merge vendor
    # =====================================================

    @staticmethod
    def merge_vendor(
        vendor_map,
        vendor
    ):

        name = str(
            vendor.get(
                "name",
                ""
            )
        ).strip()

        if not name:
            return

        if name not in vendor_map:

            vendor_map[name] = {
                "name":
                    name,

                "phone":
                    vendor.get(
                        "phone",
                        ""
                    ),

                "fax":
                    vendor.get(
                        "fax",
                        ""
                    ),

                "address":
                    vendor.get(
                        "address",
                        ""
                    ),

                "categories":
                    vendor.get(
                        "categories",
                        []
                    ),
            }

            return

        existing = vendor_map[
            name
        ]

        for field in [
            "phone",
            "fax",
            "address",
        ]:

            value = str(
                vendor.get(
                    field,
                    ""
                )
            ).strip()

            if value:

                existing[
                    field
                ] = value

    # =====================================================
    # Excel import
    # =====================================================

    @staticmethod
    def import_from_excel(file_path):

        old_vendors = (
            VendorMasterService
            .load_vendors()
        )

        vendor_map = {}

        for vendor in old_vendors:

            VendorMasterService.merge_vendor(
                vendor_map,
                vendor
            )

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        for sheet in workbook.worksheets:

            current_vendor = None

            for row in sheet.iter_rows():

                values = [
                    VendorMasterService
                    .clean_text(
                        cell.value
                    )
                    for cell in row
                ]

                # 整列合併成文字
                combined = "    ".join(
                    value
                    for value in values
                    if value
                )

                if not combined:
                    continue

                seller = (
                    VendorMasterService
                    .extract_field(
                        combined,
                        [
                            "賣方",
                            "卖方",
                        ]
                    )
                )

                if seller:

                    current_vendor = {
                        "name":
                            seller,

                        "phone":
                            VendorMasterService
                            .extract_field(
                                combined,
                                [
                                    "電話",
                                    "电话",
                                    "TEL",
                                    "Tel",
                                ]
                            ),

                        "fax":
                            VendorMasterService
                            .extract_field(
                                combined,
                                [
                                    "傳真",
                                    "传真",
                                    "FAX",
                                    "Fax",
                                ]
                            ),

                        "address":
                            VendorMasterService
                            .extract_field(
                                combined,
                                [
                                    "地址",
                                    "公司地址",
                                    "工廠地址",
                                    "厂址",
                                ]
                            ),

                        "categories":
                            [],
                    }

                    # 保留舊分類
                    if seller in vendor_map:

                        current_vendor[
                            "categories"
                        ] = vendor_map[
                            seller
                        ].get(
                            "categories",
                            []
                        )

                    VendorMasterService.merge_vendor(
                        vendor_map,
                        current_vendor
                    )

                    continue

                # 若後面的地址分另一列
                if current_vendor:

                    phone = (
                        VendorMasterService
                        .extract_field(
                            combined,
                            [
                                "電話",
                                "电话",
                                "TEL",
                                "Tel",
                            ]
                        )
                    )

                    fax = (
                        VendorMasterService
                        .extract_field(
                            combined,
                            [
                                "傳真",
                                "传真",
                                "FAX",
                                "Fax",
                            ]
                        )
                    )

                    address = (
                        VendorMasterService
                        .extract_field(
                            combined,
                            [
                                "地址",
                                "公司地址",
                                "工廠地址",
                                "厂址",
                            ]
                        )
                    )

                    if phone:
                        current_vendor[
                            "phone"
                        ] = phone

                    if fax:
                        current_vendor[
                            "fax"
                        ] = fax

                    if address:
                        current_vendor[
                            "address"
                        ] = address

                    VendorMasterService.merge_vendor(
                        vendor_map,
                        current_vendor
                    )

        vendors = list(
            vendor_map.values()
        )

        vendors.sort(
            key=lambda vendor:
            str(
                vendor.get(
                    "name",
                    ""
                )
            )
        )

        VendorMasterService.save_vendors(
            vendors
        )

        return vendors