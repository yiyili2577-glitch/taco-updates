import warnings

import pandas as pd
from openpyxl import load_workbook


# 隱藏正航 / Excel 特殊頁首頁尾警告
warnings.filterwarnings(
    "ignore",
    message="Cannot parse header or footer"
)


class ExcelService:

    PRODUCT_NO_ALIASES = [
        "產品編號",
        "料號",
        "產品編碼",
        "品號",
    ]

    PRODUCT_NAME_ALIASES = [
        "產品名稱",
        "品名",
        "產品名",
    ]

    # =====================================================
    # 清理文字
    # =====================================================

    @staticmethod
    def clean_text(value):

        if value is None:
            return ""

        return (
            str(value)
            .replace("\n", "")
            .replace("\r", "")
            .replace("\u3000", "")
            .strip()
        )

    # =====================================================
    # 判斷某列是不是產品表頭
    # =====================================================

    @staticmethod
    def is_product_header(row_values):

        clean_values = [
            ExcelService.clean_text(value)
            for value in row_values
        ]

        has_product_no = any(
            alias in clean_values
            for alias in ExcelService.PRODUCT_NO_ALIASES
        )

        has_product_name = any(
            alias in clean_values
            for alias in ExcelService.PRODUCT_NAME_ALIASES
        )

        return (
            has_product_no
            and
            has_product_name
        )

    # =====================================================
    # 找產品表頭在哪一列
    # =====================================================

    @staticmethod
    def find_header_row(
        file_path,
        sheet_name=None
    ):

        workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=False
        )

        if sheet_name:

            sheets = [
                workbook[
                    sheet_name
                ]
            ]

        else:

            sheets = (
                workbook.worksheets
            )

        for sheet in sheets:

            # 前 100 列內尋找
            max_rows = min(
                sheet.max_row,
                100
            )

            for row_index in range(
                1,
                max_rows + 1
            ):

                values = [
                    sheet.cell(
                        row_index,
                        column_index
                    ).value
                    for column_index
                    in range(
                        1,
                        sheet.max_column + 1
                    )
                ]

                if (
                    ExcelService
                    .is_product_header(
                        values
                    )
                ):

                    # pandas header 是 0-based
                    return (
                        sheet.title,
                        row_index - 1
                    )

        return (
            None,
            None
        )

    # =====================================================
    # 一般 Excel 讀取
    #
    # 先嘗試尋找產品表頭
    # 找不到才回退到一般 pandas 讀法
    # =====================================================

    @staticmethod
    def read_excel(file_path):

        try:

            sheet_name, header_row = (
                ExcelService
                .find_header_row(
                    file_path
                )
            )

            # =============================================
            # 找到產品表頭
            # =============================================

            if (
                sheet_name is not None
                and
                header_row is not None
            ):

                df = pd.read_excel(
                    file_path,
                    sheet_name=sheet_name,
                    header=header_row,
                    engine="openpyxl"
                )

            # =============================================
            # 找不到產品表頭
            # 可能是庫存分析類 Excel
            # =============================================

            else:

                df = pd.read_excel(
                    file_path,
                    engine="openpyxl"
                )

            # =============================================
            # 清理欄位
            # =============================================

            clean_columns = []

            for column in df.columns:

                clean_columns.append(
                    ExcelService
                    .clean_text(
                        column
                    )
                )

            df.columns = (
                clean_columns
            )

            # 移除完全空白列
            df = df.dropna(
                how="all"
            )

            return df.reset_index(
                drop=True
            )

        except Exception as e:

            raise Exception(
                f"Excel 讀取失敗：{str(e)}"
            )

    # =====================================================
    # 專門讀產品 Excel
    # 可以讀訂購單
    # =====================================================

    @staticmethod
    def read_product_excel(
        file_path
    ):

        sheet_name, header_row = (
            ExcelService
            .find_header_row(
                file_path
            )
        )

        if (
            sheet_name is None
            or
            header_row is None
        ):

            raise Exception(
                "找不到產品資料表頭。\n"
                "Excel 必須包含「產品編號＋品名」"
                "或「料號＋品名」。"
            )

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            header=header_row,
            engine="openpyxl"
        )

        df.columns = [
            ExcelService.clean_text(
                column
            )
            for column in df.columns
        ]

        df = df.dropna(
            how="all"
        )

        return df.reset_index(
            drop=True
        )

    # =====================================================
    # 多檔案批次讀取產品
    # =====================================================

    @staticmethod
    def read_multiple_product_excels(
        file_paths
    ):

        frames = []

        errors = []

        for file_path in file_paths:

            try:

                df = (
                    ExcelService
                    .read_product_excel(
                        file_path
                    )
                )

                df[
                    "_來源檔案"
                ] = file_path

                frames.append(
                    df
                )

            except Exception as e:

                errors.append(
                    (
                        file_path,
                        str(e)
                    )
                )

        if not frames:

            error_text = "\n".join(
                f"{file_path}\n{error}"
                for file_path, error
                in errors
            )

            raise Exception(
                "選取的 Excel 都沒有成功讀到產品。\n\n"
                + error_text
            )

        combined = pd.concat(
            frames,
            ignore_index=True
        )

        return (
            combined,
            errors
        )