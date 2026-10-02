import warnings

import pandas as pd
from openpyxl import load_workbook


class ProductExcelService:

    # =====================================================
    # 清理文字
    # =====================================================

    @staticmethod
    def clean_text(value):

        if value is None:
            return ""

        try:
            if pd.isna(value):
                return ""
        except Exception:
            pass

        return str(value).strip()

    # =====================================================
    # 尋找真正產品表頭
    #
    # 你的訂購單通常：
    #
    # 第1列 公司名稱
    # 第2列 訂單號碼
    # 第3列 日期
    # 第4列 買方
    # 第5列 賣方
    # 第7列 產品編號/品名/規格...
    # =====================================================

    @staticmethod
    def find_header_row(
        worksheet
    ):

        for row_index in range(
            1,
            worksheet.max_row + 1
        ):

            values = [
                ProductExcelService
                .clean_text(
                    worksheet.cell(
                        row=row_index,
                        column=column_index
                    ).value
                )
                for column_index
                in range(
                    1,
                    worksheet.max_column + 1
                )
            ]

            normalized = [
                value.replace(
                    " ",
                    ""
                )
                for value in values
            ]

            has_product_no = (
                "產品編號"
                in normalized
                or
                "料號"
                in normalized
                or
                "產品編碼"
                in normalized
            )

            has_name = (
                "品名"
                in normalized
                or
                "產品名稱"
                in normalized
            )

            if (
                has_product_no
                and
                has_name
            ):

                return row_index

        return None

    # =====================================================
    # 欄位名稱標準化
    # =====================================================

    @staticmethod
    def normalize_header(
        value
    ):

        text = (
            ProductExcelService
            .clean_text(
                value
            )
            .replace("\n", "")
            .replace("\r", "")
            .replace("\u3000", "")
            .replace(" ", "")
        )

        aliases = {
            "料號":
                "產品編號",

            "產品編碼":
                "產品編號",

            "產品編號":
                "產品編號",

            "品名":
                "產品名稱",

            "產品名稱":
                "產品名稱",

            "規格":
                "規格",

            "包裝":
                "包裝方式",

            "包裝方式":
                "包裝方式",
        }

        return aliases.get(
            text,
            text
        )

    # =====================================================
    # 讀取一份訂購單
    # =====================================================

    @staticmethod
    def read_products(
        file_path
    ):

        # 隱藏 openpyxl 頁首頁尾 warning
        with warnings.catch_warnings():

            warnings.filterwarnings(
                "ignore",
                message=(
                    "Cannot parse header or footer"
                )
            )

            workbook = load_workbook(
                file_path,
                data_only=True
            )

        all_products = []

        for worksheet in (
            workbook.worksheets
        ):

            header_row = (
                ProductExcelService
                .find_header_row(
                    worksheet
                )
            )

            if header_row is None:
                continue

            # =============================================
            # 建立欄位索引
            # =============================================

            column_map = {}

            for column_index in range(
                1,
                worksheet.max_column + 1
            ):

                header_value = (
                    worksheet.cell(
                        row=header_row,
                        column=column_index
                    ).value
                )

                header_name = (
                    ProductExcelService
                    .normalize_header(
                        header_value
                    )
                )

                if header_name:

                    column_map[
                        header_name
                    ] = column_index

            if (
                "產品編號"
                not in column_map
            ):

                continue

            # =============================================
            # 從表頭下一列開始讀
            # =============================================

            for row_index in range(
                header_row + 1,
                worksheet.max_row + 1
            ):

                product_no = (
                    ProductExcelService
                    .clean_text(
                        worksheet.cell(
                            row=row_index,
                            column=(
                                column_map[
                                    "產品編號"
                                ]
                            )
                        ).value
                    )
                )

                # -----------------------------------------
                # 沒產品編號就不算正式產品
                # -----------------------------------------

                if not product_no:
                    continue

                # 排除合計、備註等
                if product_no in [
                    "合計",
                    "共",
                    "備註",
                ]:

                    continue

                # -----------------------------------------
                # 名稱
                # -----------------------------------------

                product_name = ""

                if (
                    "產品名稱"
                    in column_map
                ):

                    product_name = (
                        ProductExcelService
                        .clean_text(
                            worksheet.cell(
                                row=row_index,
                                column=(
                                    column_map[
                                        "產品名稱"
                                    ]
                                )
                            ).value
                        )
                    )

                # -----------------------------------------
                # 規格
                # -----------------------------------------

                specification = ""

                if (
                    "規格"
                    in column_map
                ):

                    specification = (
                        ProductExcelService
                        .clean_text(
                            worksheet.cell(
                                row=row_index,
                                column=(
                                    column_map[
                                        "規格"
                                    ]
                                )
                            ).value
                        )
                    )

                # -----------------------------------------
                # 包裝方式
                # -----------------------------------------

                packing = ""

                if (
                    "包裝方式"
                    in column_map
                ):

                    packing = (
                        ProductExcelService
                        .clean_text(
                            worksheet.cell(
                                row=row_index,
                                column=(
                                    column_map[
                                        "包裝方式"
                                    ]
                                )
                            ).value
                        )
                    )

                all_products.append(
                    {
                        "產品編號":
                            product_no,

                        "產品名稱":
                            product_name,

                        "規格":
                            specification,

                        "包裝方式":
                            packing,
                    }
                )

        if not all_products:

            raise Exception(
                "找不到產品資料。\n\n"
                "程式會自動尋找包含「產品編號」與「品名」的表頭列。"
            )

        # =================================================
        # 同一份 Excel 內去重
        # =================================================

        product_map = {}

        for product in (
            all_products
        ):

            product_no = product[
                "產品編號"
            ]

            product_map[
                product_no
            ] = product

        products = list(
            product_map.values()
        )

        products.sort(
            key=lambda item:
            item[
                "產品編號"
            ]
        )

        return products

    # =====================================================
    # 一次讀多份 Excel
    # =====================================================

    @staticmethod
    def read_multiple_files(
        file_paths
    ):

        product_map = {}

        failed_files = []

        for file_path in file_paths:

            try:

                products = (
                    ProductExcelService
                    .read_products(
                        file_path
                    )
                )

                for product in products:

                    product_map[
                        product[
                            "產品編號"
                        ]
                    ] = product

            except Exception as e:

                failed_files.append(
                    (
                        file_path,
                        str(e)
                    )
                )

        products = list(
            product_map.values()
        )

        products.sort(
            key=lambda item:
            item[
                "產品編號"
            ]
        )

        return (
            products,
            failed_files
        )